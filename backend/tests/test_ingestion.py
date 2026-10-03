"""Offline tests for Task 4. No API keys, DB, or Redis needed.
Run from backend/:  python tests/test_ingestion.py
(Real Tesseract is used for OCR tests if installed; otherwise they are skipped.)
"""
import asyncio
import base64
import io
import os
import shutil
import sys
import tempfile
from datetime import datetime
from pathlib import Path
from types import SimpleNamespace

os.environ.setdefault("POSTGRES_PASSWORD", "x")
os.environ.setdefault("REDIS_PASSWORD", "x")
BACKEND = Path(__file__).resolve().parent.parent
sys.path[:0] = [str(BACKEND), str(BACKEND.parent)]

import httpx  # noqa: E402
from docx import Document  # noqa: E402
from fastapi import FastAPI  # noqa: E402
from fastapi.testclient import TestClient  # noqa: E402
from PIL import Image, ImageDraw, ImageFont  # noqa: E402

from api.ingest import router  # noqa: E402
from config import settings  # noqa: E402
from pipelines.ingestion import chunker  # noqa: E402
from pipelines.ingestion.chunker import chunk_pages, select_strategy, split_sentences  # noqa: E402
from pipelines.ingestion.extractors import (ExtractedPage, ExtractionError, extract_csv, extract_docx,  # noqa: E402
                                            extract_image, extract_pdf, extract_url)
from pipelines.ingestion.extractors import pdf_extractor  # noqa: E402
from pipelines.ingestion.extractors._common import rows_to_markdown  # noqa: E402
from pipelines.ingestion.pipeline import IngestionPipeline  # noqa: E402
from providers.base import GenerationResult  # noqa: E402

TMP = Path(tempfile.mkdtemp())
HAS_TESSERACT = shutil.which("tesseract") is not None


# ----------------------------------------------------------------------------- helpers
def make_pdf(texts):
    n = len(texts)
    kids = " ".join(f"{4 + 2 * i} 0 R" for i in range(n))
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            f"<< /Type /Pages /Kids [{kids}] /Count {n} >>".encode(),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    for i, t in enumerate(texts):
        stream = f"BT /F1 12 Tf 72 720 Td ({t}) Tj ET".encode()
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {5 + 2 * i} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> >> >>".encode())
        objs.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out))
        out += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objs) + 1}\n".encode() + b"0000000000 65535 f \n"
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offs)
    out += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    return out


def text_image(text, size=(1200, 400), font_size=120):
    img = Image.new("RGB", size, "white")
    ImageDraw.Draw(img).text((40, 80), text, fill="black", font=ImageFont.load_default(size=font_size))
    return img


class FakeClient:
    def __init__(self):
        self.chat_calls, self.embed_sizes = [], []

    async def chat(self, messages, routing_criteria=None, **kw):
        self.chat_calls.append((messages, routing_criteria))
        return GenerationResult("A white page with large black text.", "gpt-4o-mini", 10, 5, 1.0, 0.0, "stop")

    async def embed(self, texts):
        self.embed_sizes.append(len(texts))
        return [[float(len(t) % 5 + 1), 0.5, 0.25, 1.0] for t in texts]


def sentences(n, prefix="Sentence"):
    return " ".join(f"{prefix} number {i} describes alpha and beta behaviour in reasonable detail." for i in range(n))


# ----------------------------------------------------------------------------- extractors
def test_csv_small_and_large():
    small = TMP / "small.csv"
    small.write_text("name,score\n" + "\n".join(f"n{i},{i}" for i in range(250)))
    pages = extract_csv(str(small))
    assert len(pages) == 3 and all(p.content_type == "table" for p in pages)  # 100+100+50 rows
    assert pages[0].content.startswith("| name | score |") and pages[2].metadata["rows"] == "201-250"

    big = TMP / "big.csv"
    big.write_text("city,pop\n" + "\n".join(f"{['A','B','C'][i % 3]},{i}" for i in range(1500)))
    pages = extract_csv(str(big))
    assert pages[0].metadata["mode"] == "summary"
    assert "pop: 0 / 1499 / 749.5" in pages[0].content and "city: " in pages[0].content
    assert pages[1].metadata["mode"] == "sample_rows"


def test_rows_to_markdown_escapes():
    md = rows_to_markdown([["a", "b"], ["x|y", None]])
    assert "x\\|y" in md and md.splitlines()[1] == "| --- | --- |"


def test_docx_hierarchy_tables_headers():
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = "ACME Confidential"
    doc.add_heading("Introduction", 1)
    doc.add_paragraph("This is the intro text. It has two sentences.")
    doc.add_heading("Background", 2)
    doc.add_paragraph("Some background material goes here.")
    t = doc.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Key", "Value"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "k1", "v1"
    doc.add_heading("Methods", 1)
    doc.add_paragraph("We did things.")
    path = TMP / "t.docx"
    doc.save(path)
    pages = extract_docx(str(path))
    by = {(p.metadata.get("section"), p.content_type): p for p in pages}
    assert any(p.metadata.get("role") == "header" for p in pages)
    assert by[("Introduction", "text")].metadata["level"] == "h1"
    bg = by[("Background", "text")].metadata
    assert bg["level"] == "h2" and bg["heading_path"] == ["Introduction", "Background"]
    assert ("Background", "table") in by and "| Key | Value |" in by[("Background", "table")].content
    assert [p.page_number for p in pages] == list(range(1, len(pages) + 1))
    return str(path)


def test_pdf_digital_and_scanned():
    digital = TMP / "d.pdf"
    digital.write_bytes(make_pdf(["First page with plenty of digital text for the extractor to find.",
                                  "Second page also has enough characters to count as digital."]))
    pages = extract_pdf(str(digital))
    texts = [p for p in pages if p.content_type == "text"]
    assert [p.page_number for p in texts] == [1, 2] and not texts[0].metadata["ocr"]
    assert "digital text" in texts[0].content

    scanned = TMP / "s.pdf"
    text_image("INVOICE 12345", (1200, 500)).save(scanned, "PDF")
    if HAS_TESSERACT:
        pages = extract_pdf(str(scanned))
        assert pages and pages[0].metadata["ocr"] is True and "INVOICE" in pages[0].content.upper()
    # OCR path with the engine stubbed (always runs)
    orig = pdf_extractor.ocr_image
    pdf_extractor.ocr_image = lambda img, cfg="": "stubbed ocr text"
    try:
        pages = extract_pdf(str(scanned))
        assert pages[0].content == "stubbed ocr text" and pages[0].metadata["ocr"] is True
    finally:
        pdf_extractor.ocr_image = orig


async def test_image_resize_vision_ocr():
    path = TMP / "big.png"
    img = text_image("HELLO WORLD", (3000, 2000), 260)
    img.save(path)
    client = FakeClient()
    pages = await extract_image(str(path), client)
    messages, criteria = client.chat_calls[0]
    assert criteria.require_vision is True
    url = messages[0].content[1]["image_url"]["url"]
    sent = Image.open(io.BytesIO(base64.b64decode(url.split(",", 1)[1])))
    assert max(sent.size) <= 1024, sent.size
    assert pages[0].content_type == "image_description" and pages[0].content.startswith("A white page")
    if HAS_TESSERACT:
        assert "Text found in image:" in pages[0].content
    bad = TMP / "x.gif"
    Image.new("RGB", (10, 10)).save(bad, "GIF")
    try:
        await extract_image(str(bad), client)
        raise AssertionError("gif should be rejected")
    except ExtractionError:
        pass


async def test_url_extractor_robots_ssrf():
    body = "".join(f"<p>Paragraph {i} explains how retrieval augmented generation systems combine search "
                   f"with language models to produce grounded answers for users everywhere.</p>" for i in range(8))
    html = ("<html><head><title>My Article</title><meta name='author' content='Jane Doe'></head>"
            f"<body><article><h1>My Article</h1>{body}</article></body></html>")

    def handler(request):
        if request.url.path == "/robots.txt":
            return httpx.Response(200, text="User-agent: *\nDisallow: /private")
        return httpx.Response(200, text=html)

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    pages = await extract_url("https://example.test/post", http_client=client, validate_host=False)
    assert pages[0].metadata["title"] == "My Article" and pages[0].metadata["author"] == "Jane Doe"
    assert "retrieval augmented generation" in pages[0].content
    try:
        await extract_url("https://example.test/private/x", http_client=client, validate_host=False)
        raise AssertionError("robots.txt should block")
    except ExtractionError as exc:
        assert "robots" in str(exc)
    for bad in ("http://127.0.0.1/", "http://169.254.169.254/latest", "http://10.0.0.5/", "ftp://x.com/a"):
        try:
            await extract_url(bad)
            raise AssertionError(f"{bad} should be refused")
        except ExtractionError:
            pass


# ----------------------------------------------------------------------------- chunker
async def test_fixed_chunker():
    page = ExtractedPage(1, sentences(300), "text", {})
    chunks, strategy = await chunk_pages([page], "pdf")
    assert strategy == "fixed_size" and len(chunks) > 3
    for c in chunks:
        assert c.token_count <= 512 * 1.1 + 8, c.token_count
        assert c.content.rstrip().endswith("."), "never split mid-sentence"
    first_sents = [s.text for s in split_sentences(chunks[1].content)]
    prev_sents = {s.text for s in split_sentences(chunks[0].content)}
    assert prev_sents & set(first_sents), "consecutive chunks should overlap"
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))


async def test_table_chunker_repeats_header():
    rows = [["id", "value"]] + [[str(i), "x" * 20] for i in range(400)]
    chunks, strategy = await chunk_pages([ExtractedPage(1, rows_to_markdown(rows), "table", {})], "csv")
    assert strategy == "fixed_size" and len(chunks) > 1
    assert all(c.content.startswith("| id | value |") for c in chunks)
    assert all(c.metadata["content_type"] == "table" for c in chunks)


async def test_semantic_chunker_splits_on_topic_shift():
    cats = " ".join(f"The cat number {i} sat on the warm sunny window ledge all afternoon long." for i in range(12))
    stocks = " ".join(f"The stock market index number {i} fell sharply during trading today again." for i in range(12))

    async def embed(texts):
        return [[1.0, 0.0] if "cat" in t else [0.0, 1.0] for t in texts]

    chunks, strategy = await chunk_pages([ExtractedPage(1, cats + " " + stocks, "text", {})], "pdf", embed,
                                         strategy="semantic")
    assert strategy == "semantic" and len(chunks) == 2
    assert "cat" in chunks[0].content and "stock" not in chunks[0].content
    assert "stock" in chunks[1].content and "cat" not in chunks[1].content


async def test_hierarchical_chunker():
    pages = extract_docx(test_docx_path)
    chunks, strategy = await chunk_pages(pages, "docx")
    assert strategy == "hierarchical"
    parents = [c for c in chunks if c.metadata.get("is_parent")]
    assert {p.metadata["section"] for p in parents} == {"Introduction", "Methods"}
    intro = next(p for p in parents if p.metadata["section"] == "Introduction")
    kids = [c for c in chunks if c.metadata.get("parent_chunk_id") == intro.id]
    assert kids and {k.id for k in kids} == set(intro.metadata["child_chunk_ids"])
    assert any(k.metadata["content_type"] == "table" for k in kids)
    assert any(k.metadata.get("section") == "Background" for k in kids)


def test_strategy_selection_all_five_types():
    t = lambda n=1: ExtractedPage(n, "x", "table", {})
    x = lambda n=1, **m: ExtractedPage(n, "text here.", "text", m)
    assert select_strategy([t()], "csv") == "fixed_size"
    assert select_strategy([x(1, level="h1"), x(2, level="h2")], "docx") == "hierarchical"
    assert select_strategy([x(1, level="body")], "docx") == "fixed_size"          # DOCX without headings
    assert select_strategy([x(51)], "pdf") == "semantic" and select_strategy([x(50)], "pdf") == "fixed_size"
    assert select_strategy([ExtractedPage(1, "d", "image_description", {})], "image") == "fixed_size"
    assert select_strategy([x(1)], "url") == "fixed_size"
    assert select_strategy([t(1), t(60)], "pdf") == "fixed_size"                   # tables always fixed_size


# ----------------------------------------------------------------------------- pipeline + API
class _Ctx:
    def __init__(self, v=None):
        self.v = v

    async def __aenter__(self):
        return self.v

    async def __aexit__(self, *a):
        return False


class FakePool:
    def __init__(self):
        self.log = []

    async def execute(self, q, *a):
        self.log.append(("execute", q, a))

    def acquire(self):
        pool = self

        class Conn:
            def transaction(self):
                return _Ctx()

            async def execute(self, q, *a):
                pool.log.append(("execute", q, a))

            async def executemany(self, q, rows):
                pool.log.append(("executemany", q, list(rows)))

        return _Ctx(Conn())


async def test_pipeline_end_to_end_csv():
    path = TMP / "pipe.csv"
    path.write_text("a,b\n" + "\n".join(f"{i},{i * 2}" for i in range(250)))
    pool, client = FakePool(), FakeClient()
    result = await IngestionPipeline(pool, client).run("doc-1", str(path), "csv")
    assert result["status"] == "complete" and result["chunks"] >= 3 and result["embedding_calls"] >= 1
    statements = [e[1] for e in pool.log if e[0] == "execute"]
    assert "status='processing'" in statements[0]
    assert any("status='complete'" in s for s in statements)
    inserted = next(e for e in pool.log if e[0] == "executemany")[2]
    assert len(inserted) == result["chunks"] and inserted[0][3].startswith("[")

    bad = await IngestionPipeline(FakePool(), client).run("doc-2", str(TMP / "missing.csv"), "csv")
    assert bad["status"] == "failed"


class FakeDocPool:
    def __init__(self):
        self.docs = {}

    async def fetchrow(self, q, *a):
        if "WHERE content_hash" in q:
            return next((d for d in self.docs.values() if d["content_hash"] == a[0]), None)
        if "WHERE id" in q:
            return self.docs.get(a[0])

    async def execute(self, q, *a):
        if q.startswith("INSERT INTO documents"):
            i, fn, st, h, meta, pid = a
            self.docs[i] = {"id": i, "filename": fn, "source_type": st, "content_hash": h, "status": "queued",
                            "chunk_count": None, "metadata": meta, "created_at": datetime.now()}
        elif "status='failed'" in q:
            self.docs[a[0]]["status"] = "failed"


class FakeArq:
    def __init__(self):
        self.jobs = []

    async def enqueue_job(self, fn, *args, **kw):
        self.jobs.append((fn, args, kw))


def test_api_dedup_validation_and_status():
    settings.upload_dir = str(TMP / "uploads")
    app = FastAPI()
    app.include_router(router)
    app.state.pool, app.state.arq = FakeDocPool(), FakeArq()
    c = TestClient(app)
    csv = ("a,b\n1,2\n3,4\n", )
    r1 = c.post("/ingest", files={"file": ("data.csv", csv[0], "text/csv")})
    assert r1.status_code == 202 and r1.json()["status"] == "queued" and r1.json()["duplicate"] is False
    doc_id = r1.json()["document_id"]
    assert len(app.state.arq.jobs) == 1
    fn, args, kw = app.state.arq.jobs[0]
    assert fn == "process_document" and args[0] == doc_id and args[2] == "csv" and kw["_queue_name"] == "queue:ingest"
    assert Path(args[1]).exists()

    r2 = c.post("/ingest", files={"file": ("renamed.csv", csv[0], "text/csv")})   # same bytes, new name
    assert r2.status_code == 200 and r2.json() == {"document_id": doc_id, "status": "queued", "duplicate": True}
    assert len(app.state.arq.jobs) == 1, "duplicate must not be re-enqueued"

    assert c.post("/ingest", files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 415
    assert c.post("/ingest", files={"file": ("fake.pdf", b"not a pdf", "application/pdf")}).status_code == 415
    assert c.post("/ingest", json={}).status_code == 400
    assert c.post("/ingest", json={"url": "http://127.0.0.1/admin"}).status_code == 422

    g = c.get(f"/documents/{doc_id}")
    assert g.status_code == 200 and g.json()["status"] == "queued" and g.json()["source_type"] == "csv"
    assert c.get("/documents/not-a-uuid").status_code == 404


# ----------------------------------------------------------------------------- runner
test_docx_path = None


async def main():
    global test_docx_path
    test_csv_small_and_large(); print("PASS csv extractor")
    test_rows_to_markdown_escapes(); print("PASS markdown helper")
    test_docx_path = test_docx_hierarchy_tables_headers(); print("PASS docx extractor")
    test_pdf_digital_and_scanned(); print("PASS pdf extractor (digital + scanned%s)" % ("" if HAS_TESSERACT else ", real OCR skipped"))
    await test_image_resize_vision_ocr(); print("PASS image extractor")
    await test_url_extractor_robots_ssrf(); print("PASS url extractor")
    await test_fixed_chunker(); print("PASS fixed_size chunker")
    await test_table_chunker_repeats_header(); print("PASS table chunking")
    await test_semantic_chunker_splits_on_topic_shift(); print("PASS semantic chunker")
    await test_hierarchical_chunker(); print("PASS hierarchical chunker")
    test_strategy_selection_all_five_types(); print("PASS strategy selection")
    await test_pipeline_end_to_end_csv(); print("PASS pipeline end-to-end")
    test_api_dedup_validation_and_status(); print("PASS API dedup/validation/status")
    print("\nAll Task 4 offline tests passed")


asyncio.run(main())
